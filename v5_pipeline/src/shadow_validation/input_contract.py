from __future__ import annotations

import csv
import math
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from types import MappingProxyType
from typing import Any, Iterable, Mapping

from openpyxl import load_workbook

from . import PROSPECTIVE_START_DATE
from .integrity import canonical_json_sha256, sha256_file


DATE_COLUMNS = ("occurrence_date", "registered_date")
POST_DECISION_FIELDS = frozenset({"action_result", "on_hold", "customer_comment"})

# This is deliberately closed. New source columns can still be retained in the
# raw parser, but they cannot enter a model or the sealed feature digest until a
# new model generation explicitly adds them.
INTAKE_FEATURE_FIELDS = (
    "barcode",
    "document_no",
    "customer",
    "line",
    "customer_equipment",
    "EMS Circle",
    "product_category",
    "process_code",
    "part_code",
    "part_name",
    "serial",
    "customer_serial",
    "P/R",
    "Vendor Name",
    "Bulk",
    "occurrence_category",
    "type",
    "symptom",
    "detail",
    "occurrence_process",
    "service_type",
    "material",
    "Area",
    "Maker",
    "customer_process",
    "Chamber",
    "intake_count",
    "service_count",
    "treatment_count",
    "usage_hours",
    "cumulative_usage_hours",
    "usage_cycles",
    "quantity",
    "occurrence_date",
    "registered_date",
)


class InputContractError(RuntimeError):
    """Raised before inference when an uploaded file violates the frozen contract."""


@dataclass(frozen=True, slots=True)
class InputRow:
    row_id: str
    identity_basis: str
    feature_sha256: str
    date: str
    month: str
    run_scope: str
    source_file: str
    source_row: int
    product_key: str
    raw_customer: str
    raw: Mapping[str, Any]
    feature_payload: Mapping[str, str]


@dataclass(frozen=True, slots=True)
class InputBatch:
    rows: tuple[InputRow, ...]
    input_files: tuple[Mapping[str, Any], ...]
    mode: str


def canonical_field_name(value: Any) -> str:
    normalized = unicodedata.normalize("NFKC", str(value or "")).lower()
    return "".join(character for character in normalized if character.isalnum())


_FIELD_BY_CANONICAL = {
    canonical_field_name(name): name
    for name in (*INTAKE_FEATURE_FIELDS, "action_result", "on_hold", "on hold", "customer_comment", "반출일")
}


def load_input_files(paths: Iterable[str | Path], *, mode: str) -> InputBatch:
    if mode not in {"prospective", "dry_run", "audit"}:
        raise InputContractError(f"지원하지 않는 입력 모드입니다: {mode}")
    resolved = tuple(Path(path).expanduser().resolve() for path in paths)
    if not resolved:
        raise InputContractError("입력 XLSX/CSV 파일이 없습니다")

    input_files: list[Mapping[str, Any]] = []
    pending_rows: list[
        tuple[Path, int, dict[str, Any], str, dict[str, str], str, str]
    ] = []
    for path in sorted(resolved, key=lambda item: item.name):
        if not path.is_file():
            raise InputContractError(f"입력 파일이 없습니다: {path}")
        suffix = path.suffix.lower()
        try:
            raw_rows = _read_xlsx(path) if suffix in {".xlsx", ".xlsm"} else _read_csv(path)
        except InputContractError:
            raise
        except Exception as error:
            raise InputContractError(f"파일을 읽을 수 없습니다: {path.name}: {error}") from error
        if suffix not in {".xlsx", ".xlsm", ".csv"}:
            raise InputContractError(f"지원하지 않는 파일 형식입니다: {path.name}")
        file_sha = sha256_file(path)
        input_files.append(
            MappingProxyType(
                {
                    "fileName": path.name,
                    "sha256": file_sha,
                    "sizeBytes": path.stat().st_size,
                    "rowCount": len(raw_rows),
                }
            )
        )
        for source_row, raw in raw_rows:
            normalized = _canonicalize_row(raw)
            post_values = _nonblank_postdecision_values(normalized)
            if mode == "prospective" and post_values:
                fields = ", ".join(sorted(post_values))
                raise InputContractError(
                    f"미래 예측 입력에 nonblank 사후 필드가 있습니다: {path.name} {source_row}행 ({fields})"
                )
            observed_date = _source_date(normalized)
            if not observed_date:
                raise InputContractError(
                    f"발생일/등록일을 해석할 수 없습니다: {path.name} {source_row}행"
                )
            if mode == "prospective" and observed_date < PROSPECTIVE_START_DATE:
                raise InputContractError(
                    "실제 shadow 시작일 2026-08-01 이전 행은 prospective 실행에 "
                    f"포함할 수 없습니다: {path.name} {source_row}행 ({observed_date})"
                )
            feature_payload = _feature_payload(normalized)
            clean = feature_payload.get("process_code", "")
            part = feature_payload.get("part_code", "")
            product_key = f"{clean}-{part}" if clean and part else clean or part or "UNCLASSIFIED"
            identity_basis = _identity_basis(feature_payload, product_key, observed_date)
            pending_rows.append(
                (
                    path,
                    source_row,
                    normalized,
                    observed_date,
                    feature_payload,
                    product_key,
                    identity_basis,
                )
            )

    identity_counts = Counter(item[6] for item in pending_rows)

    def _build_input_row(item: tuple[Path, int, dict[str, Any], str, dict[str, str], str, str]) -> InputRow:
        path, source_row, normalized, observed_date, feature_payload, product_key, identity_basis = item
        identity_payload: dict[str, Any] = {"identity": identity_basis}
        if identity_counts[identity_basis] > 1:
            identity_payload["sourceRow"] = source_row
        row_id = canonical_json_sha256(identity_payload)
        return InputRow(
            row_id=row_id,
            identity_basis=identity_basis,
            feature_sha256=canonical_json_sha256(feature_payload),
            date=observed_date,
            month=observed_date[:7],
            run_scope=(
                "prospective_shadow"
                if observed_date >= PROSPECTIVE_START_DATE
                else "dry_run_only"
            ),
            source_file=path.name,
            source_row=source_row,
            product_key=product_key,
            raw_customer=feature_payload.get("customer", "") or "UNCLASSIFIED",
            raw=MappingProxyType(normalized),
            feature_payload=MappingProxyType(feature_payload),
        )

    from concurrent.futures import ThreadPoolExecutor

    with ThreadPoolExecutor(max_workers=16) as pool:
        rows = list(pool.map(_build_input_row, pending_rows))

    seen_row_ids: set[str] = set()
    for row in rows:
        if row.row_id in seen_row_ids:
            raise InputContractError(f"중복 shadow row ID입니다: {row.row_id}")
        seen_row_ids.add(row.row_id)
    if not rows:
        raise InputContractError("분석 가능한 입력 행이 없습니다")
    return InputBatch(rows=tuple(rows), input_files=tuple(input_files), mode=mode)


def _read_xlsx(path: Path) -> list[tuple[int, dict[str, Any]]]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise InputContractError(f"XLSX를 읽을 수 없습니다(DRM/손상 여부 확인): {path.name}") from error
    try:
        sheet = workbook[workbook.sheetnames[0]]
        iterator = sheet.iter_rows(values_only=True)
        try:
            header_values = next(iterator)
        except StopIteration as error:
            raise InputContractError(f"첫 번째 시트가 비어 있습니다: {path.name}") from error
        headers = _headers(header_values, path.name)
        output: list[tuple[int, dict[str, Any]]] = []
        for source_row, values in enumerate(iterator, start=2):
            if all(_scalar_text(value) == "" for value in values):
                continue
            output.append((source_row, dict(zip(headers, values, strict=False))))
        return output
    finally:
        workbook.close()


def _read_csv(path: Path) -> list[tuple[int, dict[str, Any]]]:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "cp949"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                reader = csv.DictReader(handle)
                if reader.fieldnames is None:
                    raise InputContractError(f"CSV 헤더가 없습니다: {path.name}")
                headers = _headers(reader.fieldnames, path.name)
                output: list[tuple[int, dict[str, Any]]] = []
                for source_row, values in enumerate(reader, start=2):
                    mapped = {
                        headers[index]: values.get(original, "")
                        for index, original in enumerate(reader.fieldnames)
                    }
                    if any(_scalar_text(value) for value in mapped.values()):
                        output.append((source_row, mapped))
                return output
        except UnicodeDecodeError as error:
            last_error = error
    raise InputContractError(f"CSV 인코딩을 해석할 수 없습니다: {path.name}") from last_error


def _headers(values: Iterable[Any], file_name: str) -> list[str]:
    headers: list[str] = []
    seen: set[str] = set()
    for index, value in enumerate(values, start=1):
        raw = unicodedata.normalize("NFKC", str(value or "")).strip()
        if not raw:
            raw = f"__unnamed_{index}"
        canonical = canonical_field_name(raw)
        normalized = _FIELD_BY_CANONICAL.get(canonical, raw)
        if normalized in seen:
            raise InputContractError(f"정규화 후 중복 헤더입니다: {file_name} ({normalized})")
        seen.add(normalized)
        headers.append(normalized)
    return headers


def _canonicalize_row(raw: Mapping[str, Any]) -> dict[str, Any]:
    output: dict[str, Any] = {}
    for key, value in raw.items():
        canonical = canonical_field_name(key)
        normalized = _FIELD_BY_CANONICAL.get(canonical, str(key))
        output[normalized] = value
    return output


def _nonblank_postdecision_values(row: Mapping[str, Any]) -> dict[str, str]:
    output: dict[str, str] = {}
    for key, value in row.items():
        canonical = canonical_field_name(key)
        if canonical in POST_DECISION_FIELDS and _scalar_text(value):
            output[_FIELD_BY_CANONICAL.get(canonical, key)] = _scalar_text(value)
    return output


def _feature_payload(row: Mapping[str, Any]) -> dict[str, str]:
    payload: dict[str, str] = {}
    for field in INTAKE_FEATURE_FIELDS:
        value = row.get(field, "")
        if field in DATE_COLUMNS:
            payload[field] = _parse_date(value)
        else:
            payload[field] = _scalar_text(value)
    return payload


def _source_date(row: Mapping[str, Any]) -> str:
    for field in DATE_COLUMNS:
        parsed = _parse_date(row.get(field))
        if parsed:
            return parsed
    return ""


def _identity_basis(features: Mapping[str, str], product_key: str, observed_date: str) -> str:
    identity_parts = [
        features.get("barcode", ""),
        features.get("document_no", ""),
        features.get("serial", ""),
        features.get("customer_serial", ""),
        product_key,
        observed_date,
    ]
    return "\x1f".join(identity_parts)


def _scalar_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        return format(value, ".15g")
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", str(value))).strip()


def _parse_date(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        serial = float(value)
        if 20_000 <= serial <= 100_000:
            return (datetime(1899, 12, 30) + timedelta(days=serial)).date().isoformat()
    text = _scalar_text(value)
    if re.fullmatch(r"\d+(?:\.\d+)?", text):
        serial = float(text)
        if 20_000 <= serial <= 100_000:
            return (datetime(1899, 12, 30) + timedelta(days=serial)).date().isoformat()
    normalized = text.replace("년", "-").replace("월", "-").replace("일", "")
    normalized = re.sub(r"[./]", "-", normalized)
    normalized = re.sub(r"\s+", "", normalized)
    match = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", normalized)
    if not match:
        return ""
    try:
        return date(int(match.group(1)), int(match.group(2)), int(match.group(3))).isoformat()
    except ValueError:
        return ""
